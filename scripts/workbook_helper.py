import json
import os
import stat
import sys
import time
from tempfile import NamedTemporaryFile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def main() -> None:
    configure_stdio()

    if len(sys.argv) != 2:
        raise SystemExit("Usage: workbook_helper.py <info|inspect|read-batch|write-batch|write-check>")

    command = sys.argv[1]
    payload = read_json_from_stdin()

    if command == "info":
        print_json(get_workbook_info(payload))
        return

    if command == "inspect":
        print_json(inspect_workbook(payload))
        return

    if command == "read-batch":
        print_json(read_batch(payload))
        return

    if command == "write-batch":
        print_json(write_batch(payload))
        return

    if command == "write-check":
        print_json(check_write_ready(payload))
        return

    raise SystemExit(f"Unknown command: {command}")


def get_workbook_info(payload: dict) -> dict:
    ensure_supported_workbook_path(Path(payload["filePath"]))
    workbook = open_workbook(payload["filePath"])
    try:
        return {
            "sheetNames": workbook.sheetnames,
            "activeSheetName": workbook.active.title,
        }
    finally:
        workbook.close()


def inspect_workbook(payload: dict) -> dict:
    ensure_supported_workbook_path(Path(payload["filePath"]))
    workbook = open_workbook(payload["filePath"])
    try:
        sheet = pick_sheet(workbook, payload.get("sheetName"))
        return {
            "sheetNames": workbook.sheetnames,
            "activeSheetName": workbook.active.title,
            "sheetAnalysis": analyze_sheet(sheet),
        }
    finally:
        workbook.close()


def read_batch(payload: dict) -> dict:
    ensure_supported_workbook_path(Path(payload["filePath"]))
    workbook = open_workbook(payload["filePath"])
    try:
        sheet = pick_sheet(workbook, payload.get("sheetName"))
        column = payload["column"]
        start_row = int(payload["startRow"])
        batch_size = int(payload["batchSize"])
        values = []

        for row in range(start_row, start_row + batch_size):
            cell_value = sheet[f"{column}{row}"].value
            values.append("" if cell_value is None else str(cell_value))

        return {"values": values}
    finally:
        workbook.close()


def write_batch(payload: dict) -> dict:
    ensure_supported_workbook_path(Path(payload["filePath"]))
    workbook = open_workbook(payload["filePath"])
    temp_path: Path | None = None
    workbook_closed = False
    try:
        sheet = pick_sheet(workbook, payload.get("sheetName"))
        column = payload["column"]
        start_row = int(payload["startRow"])
        matrix = payload["matrix"]
        start_column_index = column_letter_to_index(column)

        for row_offset, row_values in enumerate(matrix):
            for column_offset, value in enumerate(row_values):
                sheet.cell(
                    row=start_row + row_offset,
                    column=start_column_index + column_offset,
                    value=value,
                )

        target_path = Path(payload["filePath"])
        with NamedTemporaryFile(delete=False, suffix=target_path.suffix, dir=target_path.parent) as handle:
            temp_path = Path(handle.name)

        workbook.save(temp_path)
        workbook.close()
        workbook_closed = True

        replace_with_retry(temp_path, target_path)
        return {"writtenRows": len(matrix)}
    finally:
        if not workbook_closed:
            workbook.close()
        if temp_path and temp_path.exists():
            temp_path.unlink(missing_ok=True)


def check_write_ready(payload: dict) -> dict:
    path = Path(payload["filePath"])
    ensure_supported_workbook_path(path)

    if has_read_only_attribute(path):
        return {"ok": False, "code": "READ_ONLY_ATTRIBUTE"}

    if not os.access(path, os.W_OK):
        return {"ok": False, "code": "NO_WRITE_ACCESS"}

    try:
        with path.open("r+b"):
            pass
    except PermissionError:
        return {"ok": False, "code": "LOCKED_FOR_UPDATE"}
    except OSError as error:
        return {"ok": False, "code": "WRITE_CHECK_FAILED", "detail": str(error)}

    try:
        with NamedTemporaryFile(delete=True, suffix=path.suffix, dir=path.parent):
            pass
    except PermissionError:
        return {"ok": False, "code": "NO_DIRECTORY_WRITE_ACCESS"}
    except OSError as error:
        return {"ok": False, "code": "WRITE_CHECK_FAILED", "detail": str(error)}

    return {"ok": True, "code": "OK"}


def open_workbook(file_path: str):
    path = Path(file_path)
    keep_vba = path.suffix.lower() == ".xlsm"
    return load_workbook(filename=file_path, keep_vba=keep_vba)


def ensure_supported_workbook_path(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"Workbook file does not exist: {path}")
    if not path.is_file():
        raise ValueError(f"Workbook path is not a file: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Please choose an .xlsx or .xlsm workbook.")


def has_read_only_attribute(path: Path) -> bool:
    file_attributes = getattr(path.stat(), "st_file_attributes", 0)
    read_only_flag = getattr(stat, "FILE_ATTRIBUTE_READONLY", 0)
    return bool(read_only_flag and file_attributes & read_only_flag)


def pick_sheet(workbook, sheet_name: str):
    if sheet_name:
        return workbook[sheet_name]
    return workbook.active


def analyze_sheet(sheet) -> dict:
    source_matches = find_header_columns(sheet, "source")
    target_matches = find_header_columns(sheet, "target")

    header_issue = describe_header_issue(source_matches, target_matches)
    if header_issue:
        return manual_review_result(
            header_issue,
            build_row_counts(sheet, source_matches[0] if len(source_matches) == 1 else None, None),
        )

    source_column = get_column_letter(source_matches[0])
    target_column = get_column_letter(target_matches[0])
    row_counts = build_row_counts(sheet, source_matches[0], target_matches[0])
    start_row = find_start_row(sheet, source_matches[0], target_matches[0])

    if start_row is None:
        return manual_review_result(
            "未找到待处理行：从第 2 行开始，没有找到 source 有内容且 target 为空的行。",
            row_counts,
        )

    return {
        "status": "auto_detected",
        "message": f"已自动识别 source={source_column}，target={target_column}，开始行={start_row}。",
        "sourceColumn": source_column,
        "targetColumn": target_column,
        "startRow": start_row,
        **row_counts,
    }


def find_header_columns(sheet, header_name: str) -> list[int]:
    matches: list[int] = []

    for column_index in range(1, sheet.max_column + 1):
        cell_value = normalize_text(sheet.cell(row=1, column=column_index).value).lower()
        if cell_value == header_name:
            matches.append(column_index)

    return matches


def describe_header_issue(source_matches: list[int], target_matches: list[int]) -> str | None:
    if not source_matches and not target_matches:
        return "未找到表头：第 1 行里没有名为 source 或 target 的列。"
    if not source_matches:
        return "未找到 source 表头：请确认第 1 行存在精确的 source 列名。"
    if not target_matches:
        return "未找到 target 表头：请确认第 1 行存在精确的 target 列名。"
    if len(source_matches) > 1:
        return "检测到多个 source 表头：请保留唯一的 source 列，或手动指定列。"
    if len(target_matches) > 1:
        return "检测到多个 target 表头：请保留唯一的 target 列，或手动指定列。"
    return None


def find_start_row(sheet, source_column_index: int, target_column_index: int) -> int | None:
    for row_index in range(2, sheet.max_row + 1):
        source_value = normalize_text(sheet.cell(row=row_index, column=source_column_index).value)
        target_value = normalize_text(sheet.cell(row=row_index, column=target_column_index).value)

        if source_value and not target_value:
            return row_index

    return None


def build_row_counts(sheet, source_column_index: int | None, target_column_index: int | None) -> dict:
    sourceRowCount = 0
    untranslatedRowCount = None if target_column_index is None else 0

    if source_column_index is None:
        return {
            "sourceRowCount": None,
            "untranslatedRowCount": untranslatedRowCount,
        }

    for row_index in range(2, sheet.max_row + 1):
        source_value = normalize_text(sheet.cell(row=row_index, column=source_column_index).value)
        if not source_value:
            continue

        sourceRowCount += 1
        if target_column_index is not None:
            target_value = normalize_text(sheet.cell(row=row_index, column=target_column_index).value)
            if not target_value:
                untranslatedRowCount += 1

    return {
        "sourceRowCount": sourceRowCount,
        "untranslatedRowCount": untranslatedRowCount,
    }


def manual_review_result(message: str, row_counts: dict | None = None) -> dict:
    return {
        "status": "needs_manual_review",
        "message": message,
        "sourceColumn": None,
        "targetColumn": None,
        "startRow": None,
        **(row_counts or {"sourceRowCount": None, "untranslatedRowCount": None}),
    }


def replace_with_retry(temp_path: Path, target_path: Path) -> None:
    delays = [0.0, 0.2, 0.5, 1.0, 2.0, 4.0]
    last_error: PermissionError | None = None

    for delay in delays:
        if delay > 0:
            time.sleep(delay)

        try:
            temp_path.replace(target_path)
            return
        except PermissionError as error:
            last_error = error

    raise PermissionError(
        "Could not replace the workbook file after multiple retries. "
        f"Make sure '{target_path}' is not open in Excel/WPS, is not in read-only mode, and is not being watched by memoQ, sync software, antivirus, or Explorer preview."
    ) from last_error


def print_json(payload: dict) -> None:
    sys.stdout.write(json.dumps(payload, ensure_ascii=False))


def configure_stdio() -> None:
    if hasattr(sys.stdin, "reconfigure"):
        sys.stdin.reconfigure(encoding="utf-8")
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8")


def read_json_from_stdin() -> dict:
    return json.loads(sys.stdin.read())


def normalize_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def column_letter_to_index(column: str) -> int:
    result = 0
    for char in column.upper():
        result = result * 26 + (ord(char) - 64)
    return result


if __name__ == "__main__":
    main()
